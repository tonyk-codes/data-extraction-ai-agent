from __future__ import annotations

import argparse
import copy
import hashlib
import json
import logging
import math
import re
import shutil
import subprocess
import sys
import tempfile
import time
from collections import Counter
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

SEARCH_START_DATE = None
SEARCH_END_DATE = None
MSG_FOLDER_PATH = None

ASSET_DIR = Path(r"C:\Scripts\海外逆向到货")
TEMPLATE_PATH = ASSET_DIR / "Template.xlsx"
WEIGHT_REFERENCE_PATH = ASSET_DIR / "Item Weight Reference.xlsx"
ATTACHMENT_ROOT = ASSET_DIR / "Attachments" / "海外逆向到货"
OUTPUT_FOLDER = "報告生成"
ATTACHMENTS_FOLDER = "Attachments"
SUPPORTED_EXTENSIONS = {".pdf", ".xlsx"}
OCR_LANGUAGES = ("eng",)
OCR_TIMEOUT_SECONDS = 900
OCR_JOBS = 2
MIN_DOCUMENT_CHARACTERS = 50
MIN_PAGE_CHARACTERS = 10
DEFAULT_WEIGHT = 0.1
MANUAL_VALUE = "需自行填寫"
LIGHT_YELLOW = "FFF2CC"
HK_TIMEZONE = ZoneInfo("Asia/Hong_Kong")
REPROCESS_EXISTING_REPORTS = True
VERBOSE_LOGGING = False
KEEP_WORK_DIRECTORY = False

FIXED_VALUES = {
    "country": "HK-Hong Kong SAR China",
    "im_or_ex": "1-Import",
    "type_": "55-Return",
    "sub_type": "12-Non CBG",
    "trade_country": "CN-China",
    "trans_via_hk": "TRANS_VIA_HK_YES",
    "final_destination": "HK",
    "coo": "CN-China",
    "bu": "OVERSEA_RETURN",
}

CURRENCY_LABELS = {
    "USD": "USD-美元", "CNY": "CNY-人民幣", "EUR": "EUR-歐元", "GBP": "GBP-英鎊",
    "HKD": "HKD-港幣", "SGD": "SGD-新加坡元", "MYR": "MYR-令吉", "THB": "THB-泰銖", "JPY": "JPY-日圓",
}

DECLARATION_COMPANIES = (
    "0021-Huawei Technologies Co., Ltd.",
    "1421-Huawei International Co. Limited",
    "1241-Huawei International Pte. Ltd.",
    "0311-Huawei Tech. Investment Co., Limited",
    "1321-Huawei Device (Hong Kong) Co., Limited",
    "5531-Sparkoo Technologies Hong Kong Co., Limited",
)

CITY_COUNTRY = {
    "CALLAO": "PE-Peru", "JAKARTA": "ID-Indonesia", "TANJUNG PRIOK": "ID-Indonesia", "SURABAYA": "ID-Indonesia",
    "SEMARANG": "ID-Indonesia", "BATAM": "ID-Indonesia", "MAKASSAR": "ID-Indonesia", "BELAWAN": "ID-Indonesia",
    "SHANGHAI": "CN-China", "SHENZHEN": "CN-China", "YANTIAN": "CN-China", "QINGDAO": "CN-China",
    "NINGBO": "CN-China", "DALIAN": "CN-China", "HONG KONG": "HK-Hong Kong SAR China", "SINGAPORE": "SG-Singapore",
    "BANGKOK": "TH-Thailand", "LAEM CHABANG": "TH-Thailand", "PORT KLANG": "MY-Malaysia",
    "TANJUNG PELEPAS": "MY-Malaysia", "ROTTERDAM": "NL-Netherlands", "HAMBURG": "DE-Germany",
    "FELIXSTOWE": "GB-United Kingdom", "LONG BEACH": "US-United States", "LOS ANGELES": "US-United States",
    "NEW YORK": "US-United States", "DUBAI": "AE-United Arab Emirates", "JEBEL ALI": "AE-United Arab Emirates",
    "TOKYO": "JP-Japan", "YOKOHAMA": "JP-Japan", "BUSAN": "KR-Korea (the Republic of)",
    "INCHEON": "KR-Korea (the Republic of)", "MANILA": "PH-Philippines", "HO CHI MINH": "VN-Viet Nam",
    "HAIPHONG": "VN-Viet Nam", "MUMBAI": "IN-India", "NHAVA SHEVA": "IN-India", "CHENNAI": "IN-India",
    "COLOMBO": "LK-Sri Lanka",
}

HEADER_TO_FIELD = {
    "req number*": "req_number", "country*": "country", "im or ex*": "im_or_ex", "type*": "type_",
    "sub type": "sub_type", "trade country*": "trade_country", "transport mode*": "transport_mode",
    "trade term*": "trade_term", "declaration company*": "declaration_company", "seller name*": "seller_name",
    "seller address*": "seller_address", "buyer name*": "buyer_name", "buyer address*": "buyer_address",
    "shipper name*": "shipper_name", "shipper address*": "shipper_address", "consignee name*": "consignee_name",
    "consignee address*": "consignee_address", "trans via hk": "trans_via_hk", "total packages*": "total_packages",
    "departure country*": "departure_country", "departure port": "departure_port", "final destination*": "final_destination",
    "truckno/vesselname": "vessel_name", "flightno/ voyageno": "voyage_no", "flightno/voyageno": "voyage_no",
    "mawb": "mbl", "hawb": "hbl", "import/export date": "import_export_date", "item code*": "item_code",
    "item quantity*": "quantity", "unit*": "unit", "unit net weight*": "unit_net_weight",
    "unit price*": "unit_price", "currency*": "currency", "c.o.o*": "coo", "box material": "box_material",
    "business type": "bu", "bu": "bu",
}

FIELD_LABELS = {
    "invoice_no": ("invoice no", "invoice number", "inv no"),
    "pl_no": ("packing list no", "packing list number", "p/l no", "pl no"),
    "vessel_name": ("vessel name", "ocean vessel", "arrival vessel", "vessel"),
    "voyage_no": ("voyage no", "voyage", "voy"),
    "flight_no": ("flight no", "flight number", "flight"),
    "departure_port": ("port of loading", "loading port", "load port", "pol", "origin port", "port of origin", "airport of departure", "place of receipt", "port of shipment"),
    "total_packages": ("total packages", "no. of packages", "number of packages", "packages"),
    "trade_term": ("trade term", "trade terms", "incoterm", "delivery term"),
    "seller_name": ("shipper name", "seller name", "shipper", "seller"),
    "seller_address": ("shipper address", "seller address", "shipper addr"),
    "buyer_name": ("bill to", "buyer name", "buyer"),
    "buyer_address": ("buyer address", "bill to address"),
}

ITEM_HEADER_ALIASES = {
    "item_code": ("item code", "item code*", "item no", "item no.", "part no", "part no.", "product code", "material", "item"),
    "quantity": ("qty", "quantity", "item quantity", "item quantity*"),
    "unit": ("unit", "unit*", "uom"),
    "unit_price": ("unit price", "unit price*", "price"),
    "unit_net_weight": ("unit net weight", "net weight", "nw", "nw kg", "packing net weight"),
    "description": ("description", "desc", "item description"),
}

ITEM_CODE_PATTERN = re.compile(r"^(?=.*\d)[A-Za-z0-9]+(?:[-_/\.][A-Za-z0-9]+)*$")
ITEM_CODE_EXCLUSIONS = {"carton", "pallet", "plywood pallet", "plywood", "wooden pallet", "wooden case", "wooden", "box", "case", "plastic", "packing material", "material", "packingno", "itemno", "total", "subtotal", "grand total", "合計", "總計"}
AIR_MASTER_PATTERN = re.compile(r"(?<!\d)(\d{3})[\s-]?(\d{8})(?!\d)")
SEA_REFERENCE_PATTERN = re.compile(r"\b(?=[A-Z0-9._/-]{6,40}\b)(?=.*\d)[A-Z]{2,7}[A-Z0-9._/-]{4,33}\b", re.IGNORECASE)
HOUSE_REFERENCE_PATTERN = re.compile(r"\b(?=[A-Z0-9._/-]{5,40}\b)(?=.*\d)[A-Z0-9][A-Z0-9._/-]{4,39}\b", re.IGNORECASE)
CONTAINER_PATTERN = re.compile(r"\b[A-Z]{4}\d{7}\b")
TRANSPORT_NOISE = ("INVOICE", "PACKING", "CONTRACT", "BOOKING", "CONTAINER", "PURCHASEORDER", "PONUMBER")
MASTER_AIR_LABELS = ("MAWB", "MASTER AIR WAYBILL")
MASTER_SEA_LABELS = ("MBL", "MASTER B/L", "MASTER BILL OF LADING", "B/L NO", "BILL OF LADING NO", "SEA WAYBILL NO")
HOUSE_AIR_LABELS = ("HAWB", "HOUSE AIR WAYBILL")
HOUSE_SEA_LABELS = ("HBL", "HOUSE B/L", "HOUSE BILL OF LADING")


def progress(message: str, level: str = "資訊") -> None:
    print(f"[{datetime.now(HK_TIMEZONE):%H:%M:%S}] [{level}] {message}", flush=True)


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def norm(value: Any) -> str:
    return " ".join(str(value or "").replace("\r", " ").replace("\n", " ").replace("\u00a0", " ").strip().split())


def norm_header(value: Any) -> str:
    return norm(value).casefold().replace("（", "(").replace("）", ")")


def number(value: Any) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        result = float(value)
        return result if math.isfinite(result) else None
    match = re.search(r"[-+]?\d+(?:\.\d+)?", str(value).replace(",", ""))
    return float(match.group()) if match else None


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1048576), b""):
            digest.update(block)
    return digest.hexdigest()


def save_json(path: Path, data: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_text(json.dumps(data, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    temporary.replace(path)


def dependency_status() -> dict[str, Any]:
    status = {"pymupdf": False, "pdfplumber": False, "openpyxl": False, "extract_msg": False, "ocrmypdf": shutil.which("ocrmypdf"), "tesseract": shutil.which("tesseract"), "languages": []}
    for module in ("pymupdf", "pdfplumber", "openpyxl", "extract_msg"):
        try:
            __import__(module)
            status[module] = True
        except Exception:
            pass
    if status["tesseract"]:
        result = subprocess.run([status["tesseract"], "--list-langs"], capture_output=True, text=True, timeout=30)
        if result.returncode == 0:
            status["languages"] = [line.strip() for line in result.stdout.splitlines()[1:] if line.strip()]
    status["missing_languages"] = [language for language in OCR_LANGUAGES if language not in status["languages"]]
    status["ocr_ready"] = bool(status["ocrmypdf"] and status["tesseract"] and not status["missing_languages"])
    return status


def read_pdf_native(path: Path) -> dict[str, Any]:
    import pymupdf
    import pdfplumber
    pages = []
    with pymupdf.open(path) as document:
        for page_number, page in enumerate(document, start=1):
            raw = page.get_text("dict", sort=True)
            words_raw = page.get_text("words", sort=True)
            blocks = []
            for block_index, block in enumerate(raw.get("blocks", [])):
                if block.get("type") != 0:
                    continue
                lines = []
                for line in block.get("lines", []):
                    text = "".join(span.get("text", "") for span in line.get("spans", []))
                    if text.strip():
                        lines.append(text)
                block_text = "\n".join(lines).strip()
                if block_text:
                    blocks.append({"block": block_index, "bbox": list(block.get("bbox", [])), "text": block_text})
            words = [{"bbox": [item[0], item[1], item[2], item[3]], "text": item[4], "block": item[5], "line": item[6], "word": item[7]} for item in words_raw]
            text = page.get_text("text", sort=True) or ""
            pages.append({"page": page_number, "width": page.rect.width, "height": page.rect.height, "text": text, "blocks": blocks, "words": words, "image_count": len(page.get_images(full=True)), "text_characters": len(text.strip()), "tables": []})
    with pdfplumber.open(path) as document:
        for page_number, page in enumerate(document.pages, start=1):
            tables = []
            for settings in ({}, {"vertical_strategy": "text", "horizontal_strategy": "text"}):
                try:
                    for table in page.extract_tables(settings) or []:
                        normalized = [["" if cell is None else str(cell) for cell in row] for row in table if row]
                        if normalized and normalized not in tables:
                            tables.append(normalized)
                except Exception:
                    pass
            pages[page_number - 1]["tables"] = tables
    for page in pages:
        page["needs_ocr"] = page["text_characters"] < MIN_PAGE_CHARACTERS and page["image_count"] > 0
        page["status"] = "native" if page["text_characters"] >= MIN_PAGE_CHARACTERS else "scanned" if page["image_count"] else "empty"
    return {"kind": "pdf", "page_count": len(pages), "pages": pages, "text_characters": sum(page["text_characters"] for page in pages)}


def pdf_quality(content: dict[str, Any]) -> int:
    text = "\n".join(page.get("text", "") for page in content.get("pages", [])).upper()
    labels = ("INVOICE", "PACKING LIST", "BILL OF LADING", "AIR WAYBILL", "PORT OF LOADING", "VESSEL", "VOYAGE", "MAWB", "HAWB", "MBL", "HBL", "UNIT PRICE", "QUANTITY")
    return sum(3 for label in labels if label in text) + len(AIR_MASTER_PATTERN.findall(text)) * 4 + len(CONTAINER_PATTERN.findall(text)) * 2 + min(len(text) // 500, 10)


def convert_pdf(path: Path, workdir: Path, dependencies: dict[str, Any]) -> dict[str, Any]:
    native = read_pdf_native(path)
    needs_ocr = native["text_characters"] < MIN_DOCUMENT_CHARACTERS or any(page["needs_ocr"] for page in native["pages"])
    metadata = {"attempted": False, "applied": False, "engine": "OCRmyPDF/Tesseract", "language": "eng"}
    selected = native
    if needs_ocr:
        metadata["attempted"] = True
        if dependencies["ocr_ready"]:
            workdir.mkdir(parents=True, exist_ok=True)
            identifier = sha256(path)[:16]
            output_pdf = workdir / f"{identifier}_ocr.pdf"
            sidecar = workdir / f"{identifier}_ocr.txt"
            command = [dependencies["ocrmypdf"], "--output-type", "pdf", "--sidecar", str(sidecar), "--jobs", str(OCR_JOBS), "--skip-text", "--rotate-pages", "--deskew", "-l", "eng", str(path), str(output_pdf)]
            try:
                result = subprocess.run(command, capture_output=True, text=True, timeout=OCR_TIMEOUT_SECONDS)
                if result.returncode == 0 and output_pdf.is_file():
                    recognized = read_pdf_native(output_pdf)
                    if pdf_quality(recognized) > pdf_quality(native) or recognized["text_characters"] > native["text_characters"] * 1.15:
                        selected = recognized
                        metadata.update({"applied": True, "output_pdf": str(output_pdf), "sidecar": str(sidecar) if sidecar.exists() else None})
                    else:
                        metadata["warning"] = "OCR_QUALITY_NOT_BETTER"
                else:
                    metadata["error"] = (result.stderr or result.stdout or "OCR_FAILED")[-2000:]
            except subprocess.TimeoutExpired:
                metadata["error"] = "OCR_TIMEOUT"
        else:
            metadata["error"] = "OCR_DEPENDENCY_MISSING"
    selected["ocr"] = metadata
    selected["needs_review"] = not any(page["text"].strip() or page["tables"] for page in selected["pages"])
    return selected


def sheet_text(sheet: dict[str, Any]) -> str:
    return "\n".join(" | ".join(norm(cell.get("value")) for cell in row.get("cells", [])) for row in sheet.get("rows", []))


def classify_sheet(sheet: dict[str, Any]) -> dict[str, Any]:
    text = (sheet.get("name", "") + "\n" + sheet_text(sheet))[:100000].casefold()
    invoice = sum(weight for phrase, weight in (("commercial invoice", 12), ("invoice no", 6), ("unit price", 5), ("total amount", 4), ("bill to", 3), ("currency", 2)) if phrase in text)
    packing = sum(weight for phrase, weight in (("packing list", 12), ("packing list no", 6), ("gross weight", 5), ("net weight", 5), ("carton", 4), ("pallet", 4), ("dimension", 3), ("package", 2)) if phrase in text)
    name = sheet.get("name", "").casefold()
    if "invoice" in name or name == "inv": invoice += 8
    if "pack" in name: packing += 8
    if invoice >= 10 and packing >= 10: sheet_type = "combined_invoice_packing_list"
    elif invoice >= packing and invoice >= 5: sheet_type = "invoice"
    elif packing > invoice and packing >= 5: sheet_type = "packing_list"
    else: sheet_type = "other"
    return {"sheet_type": sheet_type, "confidence": round(min(1.0, max(invoice, packing) / 15), 3), "needs_review": max(invoice, packing) < 5 or abs(invoice - packing) < 3, "scores": {"invoice": invoice, "packing_list": packing}}


def convert_xlsx(path: Path) -> dict[str, Any]:
    from openpyxl import load_workbook
    workbook = load_workbook(path, data_only=False, read_only=False, keep_links=False)
    cached = load_workbook(path, data_only=True, read_only=False, keep_links=False)
    sheets = []
    try:
        for worksheet in workbook.worksheets:
            cached_sheet = cached[worksheet.title]
            rows = []
            populated = 0
            for row in worksheet.iter_rows():
                cells = []
                for cell in row:
                    if cell.value is not None:
                        cached_value = cached_sheet[cell.coordinate].value
                        value = cell.value.isoformat() if hasattr(cell.value, "isoformat") else cell.value
                        cached_serialized = cached_value.isoformat() if hasattr(cached_value, "isoformat") else cached_value
                        cells.append({"row": cell.row, "column": cell.column, "coordinate": cell.coordinate, "value": value, "cached_value": cached_serialized, "data_type": cell.data_type})
                        populated += 1
                if cells:
                    rows.append({"row": row[0].row, "hidden": bool(worksheet.row_dimensions[row[0].row].hidden), "cells": cells})
            hidden_columns = [key for key, dimension in worksheet.column_dimensions.items() if dimension.hidden]
            hidden_rows = [key for key, dimension in worksheet.row_dimensions.items() if dimension.hidden]
            sheet = {"name": worksheet.title, "visibility": worksheet.sheet_state, "max_row": worksheet.max_row, "max_column": worksheet.max_column, "populated_cells": populated, "merged_ranges": [str(item) for item in worksheet.merged_cells.ranges], "hidden_columns": hidden_columns, "hidden_rows": hidden_rows, "rows": rows}
            sheet["classification"] = classify_sheet(sheet) if worksheet.sheet_state == "visible" else {"sheet_type": "hidden", "confidence": 1.0, "needs_review": False}
            sheets.append(sheet)
    finally:
        workbook.close()
        cached.close()
    return {"kind": "workbook", "format": "xlsx", "sheet_count": len(sheets), "sheets": sheets}


def document_text(document: dict[str, Any]) -> str:
    content = document["content"]
    if content["kind"] == "pdf":
        return "\n".join(page.get("text", "") for page in content["pages"])
    return "\n".join(sheet_text(sheet) for sheet in content["sheets"] if sheet.get("visibility") == "visible")


def classify_document(document: dict[str, Any]) -> dict[str, Any]:
    text = (document["source"]["filename"] + "\n" + document_text(document))[:150000].casefold()
    scores = Counter()
    groups = {
        "invoice": (("commercial invoice", 10), ("invoice no", 4), ("unit price", 4), ("total amount", 3), ("bill to", 2)),
        "packing_list": (("packing list", 10), ("packing list no", 4), ("gross weight", 4), ("net weight", 4), ("carton", 3), ("pallet", 3)),
        "air": (("air waybill", 8), ("mawb", 5), ("hawb", 5), ("airport of departure", 3), ("flight", 2)),
        "sea": (("bill of lading", 8), ("sea waybill", 8), ("mbl", 5), ("hbl", 5), ("port of loading", 3), ("vessel", 3), ("voyage", 3)),
        "arrival_notice": (("arrival notice", 12), ("arrival advice", 10)),
    }
    for category, terms in groups.items(): scores[category] = sum(weight for phrase, weight in terms if phrase in text)
    sheet_types = [sheet["classification"]["sheet_type"] for sheet in document["content"].get("sheets", []) if sheet.get("visibility") == "visible"]
    if "invoice" in sheet_types and "packing_list" in sheet_types:
        kind, score = "combined_invoice_packing_list", scores["invoice"] + scores["packing_list"] + 20
    elif scores["arrival_notice"] >= 10:
        kind, score = "arrival_notice", scores["arrival_notice"]
    else:
        category, score = scores.most_common(1)[0]
        if category == "air": kind = "air_waybill_master" if "mawb" in text or "master air waybill" in text else "air_waybill_house"
        elif category == "sea": kind = "bill_of_lading_master" if "mbl" in text or "master bill" in text or "sea waybill" in text else "bill_of_lading_house"
        else: kind = category if score else "unknown"
    values = sorted(scores.values(), reverse=True)
    second = values[1] if len(values) > 1 else 0
    return {"document_type": kind, "confidence": round(min(1.0, score / 15), 3), "needs_review": score < 5 or score - second < 3, "scores": dict(scores)}


def field_candidates(documents: list[dict[str, Any]], key: str, preferred: tuple[str, ...] = ()) -> list[dict[str, Any]]:
    output = []
    for document in sorted(documents, key=lambda item: 0 if item["classification"]["document_type"] in preferred else 1):
        text = document_text(document)
        for label in FIELD_LABELS[key]:
            escaped = re.escape(label)
            for pattern in (rf"(?im)^\s*{escaped}\s*[:#]?\s*([^\n\r|]{{1,180}})", rf"(?i)\b{escaped}\s*[:#]?\s*([^\n\r|,;]{{1,140}})"):
                match = re.search(pattern, text)
                if match:
                    value = norm(match.group(1)).strip("[] .:;")
                    if value: output.append({"value": value, "document_id": document["document_id"], "filename": document["source"]["filename"], "location": "text", "confidence": 0.85})
                    break
        if document["content"]["kind"] == "workbook":
            for sheet in document["content"]["sheets"]:
                if sheet.get("visibility") != "visible": continue
                grid = {(cell["row"], cell["column"]): cell for row in sheet["rows"] for cell in row["cells"]}
                for cell in grid.values():
                    if any(label.casefold() in norm(cell["value"]).casefold() for label in FIELD_LABELS[key]):
                        for offset in range(1, 4):
                            adjacent = grid.get((cell["row"], cell["column"] + offset))
                            if adjacent and norm(adjacent.get("cached_value") if adjacent.get("cached_value") is not None else adjacent["value"]):
                                output.append({"value": norm(adjacent.get("cached_value") if adjacent.get("cached_value") is not None else adjacent["value"]), "document_id": document["document_id"], "filename": document["source"]["filename"], "location": f"{sheet['name']}!{adjacent['coordinate']}", "confidence": 0.98})
                                break
    unique = []
    seen = set()
    for item in output:
        token = re.sub(r"\s+", "", item["value"]).casefold()
        if token not in seen:
            seen.add(token)
            unique.append(item)
    return unique


def field_result(documents: list[dict[str, Any]], key: str, preferred: tuple[str, ...] = ()) -> dict[str, Any]:
    candidates = field_candidates(documents, key, preferred)
    if not candidates: return {"value": None, "status": "unresolved", "confidence": 0.0, "candidates": []}
    return {"value": candidates[0]["value"], "status": "exact" if len(candidates) == 1 else "conflict", "confidence": candidates[0]["confidence"] if len(candidates) == 1 else 0.4, "candidates": candidates}


def clean_reference(value: Any) -> str | None:
    text = re.sub(r"\s+", "", norm(value).upper()).replace("\\", "-").strip(" .,:;#|()[]{}'\"")
    if not text or len(text) < 5 or len(text) > 40 or not re.search(r"\d", text): return None
    compact = text.replace("/", "").replace("-", "").replace("_", "").replace(".", "")
    if any(noise in compact for noise in TRANSPORT_NOISE): return None
    if CONTAINER_PATTERN.fullmatch(text): return None
    return text


def contextual_references(text: str, labels: tuple[str, ...], pattern: re.Pattern[str], air_master: bool = False) -> list[str]:
    found = []
    for label in sorted(labels, key=len, reverse=True):
        escaped = re.escape(label).replace(r"\ ", r"[\s._:/\\#=\-]*")
        matcher = re.compile(rf"(?i)(?<![A-Z0-9]){escaped}[\s._:/\\#=\-]*(?:(?:NO|NUMBER|NUM|REF|REFERENCE)[\s._:/\\#=\-]*)?([^\n\r|,;]{{3,80}})")
        for match in matcher.finditer(text):
            candidate = pattern.search(match.group(1).upper())
            if not candidate: continue
            if air_master:
                serial = candidate.group(2)
                value = f"{candidate.group(1)}-{serial}"
            else:
                value = candidate.group(0)
            value = clean_reference(value)
            if value: found.append(value)
    return list(dict.fromkeys(found))


def msg_subject_references(shipment: Path) -> dict[str, list[dict[str, Any]]]:
    result = {"master": [], "house": []}
    try:
        import extract_msg
    except ImportError:
        return result
    for path in sorted(shipment.glob("*.msg")):
        message = None
        try:
            message = extract_msg.Message(str(path))
            subject = norm(getattr(message, "subject", ""))
            for value in contextual_references(subject, MASTER_AIR_LABELS, AIR_MASTER_PATTERN, True) + contextual_references(subject, MASTER_SEA_LABELS, SEA_REFERENCE_PATTERN):
                result["master"].append({"value": value, "source": "msg_subject", "filename": path.name, "confidence": 0.99})
            for value in contextual_references(subject, HOUSE_AIR_LABELS + HOUSE_SEA_LABELS, HOUSE_REFERENCE_PATTERN):
                result["house"].append({"value": value, "source": "msg_subject", "filename": path.name, "confidence": 0.98})
        except Exception:
            pass
        finally:
            if message is not None:
                try: message.close()
                except Exception: pass
    return result


def document_transport_references(documents: list[dict[str, Any]], mode: str) -> dict[str, list[dict[str, Any]]]:
    result = {"master": [], "house": []}
    allowed = {"air_waybill_master", "air_waybill_house", "bill_of_lading_master", "bill_of_lading_house", "arrival_notice"}
    for document in documents:
        if document["classification"]["document_type"] not in allowed: continue
        text = document_text(document)
        masters = contextual_references(text, MASTER_AIR_LABELS, AIR_MASTER_PATTERN, True) if mode == "BY AIR" else contextual_references(text, MASTER_SEA_LABELS, SEA_REFERENCE_PATTERN)
        houses = contextual_references(text, HOUSE_AIR_LABELS, HOUSE_REFERENCE_PATTERN) if mode == "BY AIR" else contextual_references(text, HOUSE_SEA_LABELS, HOUSE_REFERENCE_PATTERN)
        for value in masters: result["master"].append({"value": value, "source": "transport_document", "document_id": document["document_id"], "filename": document["source"]["filename"], "confidence": 0.95})
        for value in houses: result["house"].append({"value": value, "source": "transport_document", "document_id": document["document_id"], "filename": document["source"]["filename"], "confidence": 0.95})
    return result


def detect_item_header(rows: list[dict[str, Any]]) -> tuple[int | None, dict[str, int]]:
    best = (None, {}, 0)
    priority = {"item code*": 0, "item code": 1, "item no": 2, "part no": 3, "part no.": 3, "product code": 4, "item": 5, "material": 6}
    for position, row in enumerate(rows):
        mapping = {}
        ranks = {}
        for cell in row["cells"]:
            value = re.sub(r"[^a-z0-9]+", " ", norm_header(cell["value"])).strip()
            for field, aliases in ITEM_HEADER_ALIASES.items():
                for alias in aliases:
                    if value == re.sub(r"[^a-z0-9]+", " ", alias.casefold()).strip():
                        rank = priority.get(alias.casefold(), 50)
                        if field not in mapping or rank < ranks[field]: mapping[field], ranks[field] = cell["column"], rank
        if "item_code" in mapping and "quantity" in mapping and len(mapping) > best[2]: best = (position, mapping, len(mapping))
    return best[0], best[1]


def valid_item_code(value: Any) -> bool:
    text = norm(value)
    return bool(text and text.casefold() not in ITEM_CODE_EXCLUSIONS and len(text) <= 60 and " " not in text and ITEM_CODE_PATTERN.fullmatch(text))


def extract_items(document: dict[str, Any]) -> list[dict[str, Any]]:
    if document["content"]["kind"] != "workbook": return []
    items = []
    for sheet in document["content"]["sheets"]:
        if sheet.get("visibility") != "visible": continue
        position, columns = detect_item_header(sheet["rows"])
        if position is None: continue
        source_type = sheet["classification"].get("sheet_type", document["classification"]["document_type"])
        for row in sheet["rows"][position + 1:]:
            values = {cell["column"]: cell.get("cached_value") if cell.get("cached_value") is not None else cell["value"] for cell in row["cells"]}
            first = next((norm(value) for value in values.values() if norm(value)), "")
            if first.casefold() in {"total", "subtotal", "grand total", "合計", "總計"}: break
            code = norm(values.get(columns["item_code"]))
            quantity = number(values.get(columns["quantity"]))
            if not valid_item_code(code) or quantity is None or quantity <= 0: continue
            items.append({"line_id": f"{document['document_id']}:{sheet['name']}:{row['row']}", "source_line_no": row["row"], "item_code": code, "quantity": quantity, "unit": norm(values.get(columns.get("unit"))) or "PCS", "unit_price": number(values.get(columns.get("unit_price"))), "unit_net_weight": number(values.get(columns.get("unit_net_weight"))), "description": norm(values.get(columns.get("description")))[:300], "source_type": source_type, "source_document": document["document_id"], "source_filename": document["source"]["filename"], "source_sheet": sheet["name"], "source_location": f"{sheet['name']}!Row {row['row']}", "provenance": [{"document_id": document["document_id"], "filename": document["source"]["filename"], "location": f"{sheet['name']}!Row {row['row']}", "source_type": source_type}]})
    return items


def load_weights() -> dict[str, float]:
    from openpyxl import load_workbook
    workbook = load_workbook(WEIGHT_REFERENCE_PATH, data_only=True, read_only=True)
    result = {}
    try:
        for row in workbook.active.iter_rows(min_row=2, values_only=True):
            if len(row) >= 2 and row[0] is not None:
                weight = number(row[1])
                if weight and weight > 0: result[norm(row[0]).casefold()] = weight
    finally:
        workbook.close()
    return result


def country_from_port(port: str | None) -> str | None:
    value = norm(port).upper()
    for city, country in CITY_COUNTRY.items():
        if city in value: return country
    return None


def build_shipment(source: Path, documents: list[dict[str, Any]]) -> dict[str, Any]:
    preferred = {"vessel_name": ("arrival_notice", "bill_of_lading_master"), "voyage_no": ("arrival_notice", "bill_of_lading_master"), "flight_no": ("air_waybill_master", "air_waybill_house"), "departure_port": (), "total_packages": (), "trade_term": ("invoice", "combined_invoice_packing_list"), "seller_name": ("packing_list", "invoice"), "seller_address": ("packing_list", "invoice"), "buyer_name": ("invoice", "combined_invoice_packing_list"), "buyer_address": ("invoice", "combined_invoice_packing_list"), "invoice_no": ("invoice", "combined_invoice_packing_list"), "pl_no": ("packing_list", "combined_invoice_packing_list")}
    fields = {key: field_result(documents, key, kinds) for key, kinds in preferred.items()}
    all_text = "\n".join(document_text(document) for document in documents)
    kinds = [document["classification"]["document_type"] for document in documents]
    if any(kind.startswith("air_waybill") for kind in kinds) or "by air" in all_text.casefold(): mode = "BY AIR"
    elif any(kind.startswith("bill_of_lading") or kind == "arrival_notice" for kind in kinds) or any(term in all_text.casefold() for term in ("vessel", "voyage", "port of loading", "by sea")): mode = "BY SEA"
    else: mode = "BY AIR"
    msg_refs = msg_subject_references(source)
    doc_refs = document_transport_references(documents, mode)
    masters = msg_refs["master"] + doc_refs["master"]
    houses = msg_refs["house"] + doc_refs["house"]
    mbl = masters[0]["value"] if masters else None
    hbl = houses[0]["value"] if houses else None
    invoice_no = clean_reference(fields["invoice_no"]["value"])
    pl_no = clean_reference(fields["pl_no"]["value"])
    folder_no = clean_reference(source.name)
    req_number = next((value for value in (hbl, invoice_no, pl_no, folder_no) if value), None)
    items = []
    for document in documents:
        if document["classification"]["document_type"] in {"invoice", "packing_list", "combined_invoice_packing_list"}: items.extend(extract_items(document))
    weights = load_weights()
    currency = next((CURRENCY_LABELS[code] for code in CURRENCY_LABELS if re.search(rf"\b{code}\b", all_text.upper())), CURRENCY_LABELS["USD"])
    for item in items:
        key = item["item_code"].casefold()
        if key in weights: item["unit_net_weight"], item["weight_source"] = weights[key], "item_weight_reference"
        elif item.get("unit_net_weight") and item["unit_net_weight"] > 0: item["weight_source"] = item["source_type"]
        else: item["unit_net_weight"], item["weight_source"] = DEFAULT_WEIGHT, "fallback_0.1"
        item["currency"] = currency
        item["coo"] = "CN-China"
    buyer = fields["buyer_name"]["value"]
    declaration = next((company for company in DECLARATION_COMPANIES if company.split("-", 1)[-1].casefold() in norm(buyer).casefold()), DECLARATION_COMPANIES[3])
    trade_text = norm(fields["trade_term"]["value"]).upper()
    trade_term = "DAP" if "DAP" in trade_text else "CIP"
    voyage = fields["flight_no"]["value"] if mode == "BY AIR" else fields["voyage_no"]["value"]
    if mode == "BY SEA" and voyage and not norm(voyage).upper().startswith("V."): voyage = "V." + norm(voyage)
    header = {"req_number": req_number, "transport_mode": mode, "trade_term": trade_term, "declaration_company": declaration, "seller_name": fields["seller_name"]["value"], "seller_address": fields["seller_address"]["value"], "buyer_name": buyer, "buyer_address": fields["buyer_address"]["value"], "shipper_name": fields["seller_name"]["value"], "shipper_address": fields["seller_address"]["value"], "consignee_name": "Huawei Tech. Investment Co., Limited", "consignee_address": "9th Floor, Tower 6, The Gateway, No. 9 Canton Road, Tsim Sha Tsui, Kowloon, Hong Kong", "total_packages": number(fields["total_packages"]["value"]), "departure_country": country_from_port(fields["departure_port"]["value"]), "departure_port": fields["departure_port"]["value"], "final_destination": "HK", "vessel_name": None if mode == "BY AIR" else fields["vessel_name"]["value"], "voyage_no": voyage, "mbl": mbl, "hbl": hbl, "import_export_date": MANUAL_VALUE, "bu": "OVERSEA_RETURN"}
    metadata = {**fields, "mbl": {"value": mbl, "status": "exact" if len(masters) == 1 else "conflict" if masters else "unresolved", "confidence": masters[0]["confidence"] if masters else 0.0, "candidates": masters}, "hbl": {"value": hbl, "status": "exact" if len(houses) == 1 else "conflict" if houses else "unresolved", "confidence": houses[0]["confidence"] if houses else 0.0, "candidates": houses}, "req_number": {"value": req_number, "status": "exact" if hbl or invoice_no or pl_no else "folder_fallback" if folder_no else "unresolved", "confidence": 0.98 if hbl else 0.95 if invoice_no or pl_no else 0.75 if folder_no else 0.0, "candidates": [{"type": kind, "value": value} for kind, value in (("hbl", hbl), ("invoice_no", invoice_no), ("pl_no", pl_no), ("filename", folder_no)) if value]}, "transport_mode": {"value": mode, "status": "exact" if any(kind.startswith(("air_waybill", "bill_of_lading")) or kind == "arrival_notice" for kind in kinds) else "defaulted", "confidence": 0.95 if any(kind.startswith(("air_waybill", "bill_of_lading")) or kind == "arrival_notice" for kind in kinds) else 0.4}, "trade_term": {"value": trade_term, "status": "exact" if "CIP" in trade_text or "DAP" in trade_text else "defaulted", "confidence": 0.9 if "CIP" in trade_text or "DAP" in trade_text else 0.4}, "departure_country": {"value": header["departure_country"], "status": "exact" if header["departure_country"] else "unresolved", "confidence": 0.9 if header["departure_country"] else 0.0}}
    return {"shipment_id": req_number or source.name, "group_key": f"{mbl or ''}|{hbl or ''}", "source_folder": str(source), "source_item_count": len(items), "header": header, "items": items, "field_metadata": metadata, "req_number_resolution": {"rule_version": "2.0", "selected_value": req_number, "selection_rule": "hbl_then_invoice_then_pl_then_filename", "candidates": metadata["req_number"]["candidates"]}, "documents": [{"document_id": document["document_id"], "filename": document["source"]["filename"], "classification": document["classification"]} for document in documents]}


def evidence_for_task(field: str, documents: list[dict[str, Any]], limit: int = 12) -> list[dict[str, Any]]:
    terms = {"departure_country": ("port of loading", "loading port", "airport of departure", "origin"), "departure_port": ("port of loading", "loading port", "airport of departure", "origin"), "mbl": ("mawb", "mbl", "master bill", "bill of lading", "sea waybill"), "hbl": ("hawb", "hbl", "house bill"), "seller_name": ("shipper", "seller"), "seller_address": ("shipper", "seller"), "buyer_name": ("bill to", "buyer"), "buyer_address": ("bill to", "buyer"), "trade_term": ("trade term", "incoterm", "delivery term"), "transport_mode": ("air waybill", "bill of lading", "vessel", "flight", "voyage")}.get(field, (field.replace("_", " "),))
    evidence = []
    for document in documents:
        content = document["content"]
        if content["kind"] == "pdf":
            for page in content["pages"]:
                for block in page.get("blocks", []):
                    if any(term in block["text"].casefold() for term in terms):
                        evidence.append({"document_id": document["document_id"], "filename": document["source"]["filename"], "location": f"page {page['page']}, block {block['block']}", "text": block["text"][:1000], "ocr_derived": bool(content.get("ocr", {}).get("applied"))})
        else:
            for sheet in content["sheets"]:
                for row in sheet["rows"]:
                    row_text = " | ".join(norm(cell.get("cached_value") if cell.get("cached_value") is not None else cell["value"]) for cell in row["cells"])
                    if any(term in row_text.casefold() for term in terms):
                        evidence.append({"document_id": document["document_id"], "filename": document["source"]["filename"], "location": f"{sheet['name']}!Row {row['row']}", "text": row_text[:1000], "ocr_derived": False})
        if len(evidence) >= limit: break
    return evidence[:limit]


def review_tasks(shipment: dict[str, Any], documents: list[dict[str, Any]], failures: list[dict[str, Any]]) -> list[dict[str, Any]]:
    tasks = []
    for field in ("req_number", "transport_mode", "trade_term", "declaration_company", "seller_name", "seller_address", "buyer_name", "buyer_address", "departure_country", "departure_port", "mbl", "hbl"):
        meta = shipment["field_metadata"].get(field, {"status": "unresolved", "confidence": 0.0, "candidates": []})
        if shipment["header"].get(field) in (None, "") or meta.get("status") in {"unresolved", "conflict", "defaulted"} or meta.get("confidence", 1.0) < 0.75:
            tasks.append({"task_id": f"T{len(tasks)+1:04d}", "entity_type": "shipment", "field": field, "current_value": shipment["header"].get(field), "reason_code": "MISSING_OR_UNCERTAIN_FIELD", "confidence": meta.get("confidence", 0.0), "candidates": meta.get("candidates", []), "resolution_owner": "ai_agent", "ai_editable": True, "raw_file_access_allowed": False, "evidence": evidence_for_task(field, documents)})
    for document in documents:
        if document["classification"]["needs_review"]:
            tasks.append({"task_id": f"T{len(tasks)+1:04d}", "entity_type": "document", "document_id": document["document_id"], "filename": document["source"]["filename"], "field": "document_type", "current_value": document["classification"]["document_type"], "reason_code": "LOW_CLASSIFICATION_CONFIDENCE", "confidence": document["classification"]["confidence"], "resolution_owner": "ai_agent", "ai_editable": True, "raw_file_access_allowed": False, "evidence": evidence_for_task("transport_mode", [document])})
    if not shipment["items"]:
        tasks.append({"task_id": f"T{len(tasks)+1:04d}", "entity_type": "shipment", "field": "items", "current_value": [], "reason_code": "NO_ITEM_ROWS_EXTRACTED", "resolution_owner": "ai_agent", "ai_editable": True, "raw_file_access_allowed": False, "evidence": []})
    for item in shipment["items"]:
        if item.get("unit_price") is None:
            tasks.append({"task_id": f"T{len(tasks)+1:04d}", "entity_type": "item", "line_id": item["line_id"], "item_code": item["item_code"], "field": "unit_price", "current_value": None, "reason_code": "MISSING_UNIT_PRICE", "resolution_owner": "ai_agent", "ai_editable": True, "raw_file_access_allowed": False, "evidence": item["provenance"]})
    for failure in failures:
        tasks.append({"task_id": f"T{len(tasks)+1:04d}", "entity_type": "document", "filename": failure["filename"], "field": "conversion", "current_value": None, "reason_code": "DOCUMENT_CONVERSION_FAILED", "resolution_owner": "manual", "ai_editable": False, "raw_file_access_allowed": False, "evidence": [], "details": failure["error"]})
    return tasks


def validate(shipment: dict[str, Any], tasks: list[dict[str, Any]]) -> dict[str, Any]:
    issues = []
    if not shipment["header"]["req_number"]: issues.append({"severity": "error", "code": "MISSING_REQ_NUMBER"})
    if shipment["header"]["transport_mode"] not in {"BY AIR", "BY SEA"}: issues.append({"severity": "error", "code": "INVALID_TRANSPORT_MODE"})
    if not shipment["items"]: issues.append({"severity": "error", "code": "EMPTY_ITEMS"})
    for item in shipment["items"]:
        if not valid_item_code(item["item_code"]): issues.append({"severity": "error", "code": "INVALID_ITEM_CODE", "line_id": item["line_id"]})
        if item["quantity"] <= 0: issues.append({"severity": "error", "code": "INVALID_QUANTITY", "line_id": item["line_id"]})
    status = "FAILED" if any(issue["severity"] == "error" for issue in issues) else "DRAFT_REVIEW_REQUIRED" if any(task.get("ai_editable") for task in tasks) else "FINAL"
    return {"schema_version": "1.0", "generated_at": utc_now(), "status": status, "valid_for_final": status == "FINAL", "issues": issues, "summary": {"error_count": sum(issue["severity"] == "error" for issue in issues), "review_task_count": len(tasks)}}


def report_rows(shipment: dict[str, Any]) -> list[dict[str, Any]]:
    rows = []
    for index, item in enumerate(shipment["items"]):
        header = shipment["header"]
        rows.append({**FIXED_VALUES, **header, "line_id": item["line_id"], "total_packages": header["total_packages"] if index == 0 else 0, "box_material": "110-Package" if index == 0 else "120-Ditto", "item_code": item["item_code"], "quantity": item["quantity"], "unit": item["unit"], "unit_net_weight": item["unit_net_weight"], "unit_price": item["unit_price"], "currency": item["currency"]})
    return rows


def render(report: dict[str, Any], tasks: list[dict[str, Any]], output: Path, draft: bool) -> None:
    from openpyxl import load_workbook
    from openpyxl.styles import PatternFill
    shutil.copy2(TEMPLATE_PATH, output)
    workbook = load_workbook(output)
    original_names = list(workbook.sheetnames)
    validation_counts = {name: len(workbook[name].data_validations.dataValidation) for name in original_names}
    sheet = workbook["模板"] if "模板" in workbook.sheetnames else workbook.active
    columns = {norm_header(cell.value): cell.column for cell in sheet[1] if norm_header(cell.value)}
    for row in sheet.iter_rows(min_row=2, max_row=max(sheet.max_row, 2)):
        for cell in row: cell.value = None
    yellow = PatternFill(start_color=LIGHT_YELLOW, end_color=LIGHT_YELLOW, fill_type="solid")
    ai_header_fields = {task["field"] for task in tasks if task.get("entity_type") == "shipment" and task.get("ai_editable")}
    ai_item_fields = {(task.get("line_id"), task["field"]) for task in tasks if task.get("entity_type") == "item" and task.get("ai_editable")}
    for row_number, data in enumerate(report_rows(report["shipments"][0]), start=2):
        if row_number != 2:
            for column in range(1, sheet.max_column + 1):
                sheet.cell(row_number, column)._style = copy.copy(sheet.cell(2, column)._style)
                sheet.cell(row_number, column).number_format = sheet.cell(2, column).number_format
        for header, column in columns.items():
            field = HEADER_TO_FIELD.get(header)
            if not field: continue
            value = data.get(field)
            ai_uncertain = field in ai_header_fields or (data["line_id"], field) in ai_item_fields
            sheet.cell(row_number, column).value = value
            if field == "import_export_date" or draft and ai_uncertain:
                sheet.cell(row_number, column).fill = yellow
    workbook.save(output)
    workbook.close()
    check = load_workbook(output)
    preserved = list(check.sheetnames) == original_names and all(len(check[name].data_validations.dataValidation) == validation_counts[name] for name in original_names)
    check.close()
    if not preserved:
        output.unlink(missing_ok=True)
        raise RuntimeError("輸出活頁簿結構或資料驗證未能完整保留")


def process_shipment(source: Path, verbose: bool, keep: bool) -> int:
    started = time.time()
    source = source.resolve()
    attachments = source / ATTACHMENTS_FOLDER
    output = source / OUTPUT_FOLDER
    output.mkdir(parents=True, exist_ok=True)
    if not attachments.is_dir(): raise FileNotFoundError(f"找不到附件資料夾：{attachments}")
    dependencies = dependency_status()
    files = sorted((path for path in attachments.rglob("*") if path.is_file() and path.suffix.lower() in SUPPORTED_EXTENSIONS), key=lambda path: str(path).casefold())
    progress(f"運輸資料夾：{source}")
    progress(f"發現附件數量：{len(files)}")
    if not files: raise FileNotFoundError(f"沒有 PDF 或 XLSX：{attachments}")
    workdir = Path(tempfile.mkdtemp(prefix="fong-step2-"))
    documents = []
    failures = []
    try:
        for index, path in enumerate(files, start=1):
            progress(f"正規化文件 {index}/{len(files)}：{path.name}")
            try:
                digest = sha256(path)
                content = convert_pdf(path, workdir / "ocr", dependencies) if path.suffix.lower() == ".pdf" else convert_xlsx(path)
                document = {"document_id": digest[:16], "source": {"filename": path.name, "path": str(path), "sha256": digest, "size": path.stat().st_size}, "content": content}
                document["classification"] = classify_document(document)
                normalized = output / "normalized" / document["document_id"] / "document.json"
                save_json(normalized, document)
                document["normalized_json"] = str(normalized)
                documents.append(document)
                progress(f"文件分類：{document['classification']['document_type']}；信心度={document['classification']['confidence']:.2f}")
            except Exception as exc:
                failures.append({"filename": path.name, "path": str(path), "error": f"{type(exc).__name__}: {exc}"})
                progress(f"文件轉換失敗：{path.name}；{type(exc).__name__}: {exc}", "錯誤")
        manifest = {"schema_version": "1.2", "generated_at": utc_now(), "source_folder": str(source), "attachments_folder": str(attachments), "discovered_count": len(files), "successful_count": len(documents), "failed_count": len(failures), "unaccounted_count": len(files) - len(documents) - len(failures), "documents": [{"document_id": document["document_id"], "filename": document["source"]["filename"], "normalized_json": document["normalized_json"], "classification": document["classification"], "page_count": document["content"].get("page_count"), "sheet_count": document["content"].get("sheet_count"), "sheet_classifications": [{"name": sheet["name"], "visibility": sheet["visibility"], "classification": sheet.get("classification")} for sheet in document["content"].get("sheets", [])], "ocr": document["content"].get("ocr")} for document in documents], "failures": failures}
        save_json(output / "document_manifest.json", manifest)
        if not documents: raise RuntimeError("所有附件轉換失敗")
        shipment = build_shipment(source, documents)
        tasks = review_tasks(shipment, documents, failures)
        validation = validate(shipment, tasks)
        report = {"schema_version": "5.1", "rule_versions": {"req_number": "2.0", "field_policies": "1.0"}, "generated_at": utc_now(), "source_folder": str(source), "manual_fields": [{"field": "import_export_date", "value": MANUAL_VALUE, "resolution_owner": "manual", "ai_editable": False}], "shipments": [shipment], "issues": failures}
        request = {"schema_version": "1.1", "generated_at": utc_now(), "job_id": f"{source.parent.name}_{source.name}", "status": "REVIEW_REQUIRED" if any(task.get("ai_editable") for task in tasks) else "NO_REVIEW_REQUIRED", "source_folder": str(source), "extracted_report": "extracted_report.json", "document_manifest": "document_manifest.json", "normalized_root": "normalized", "source_policy": "JSON_ONLY", "raw_file_access_allowed": False, "manual_fields": [{"field": "import_export_date", "value": MANUAL_VALUE, "ai_editable": False}], "review_tasks": tasks}
        save_json(output / "extracted_report.json", report)
        save_json(output / "validation.json", validation)
        save_json(output / "ai_review_request.json", request)
        report_path = None
        if validation["status"] != "FAILED":
            report_path = output / ("Final_Report.xlsx" if validation["status"] == "FINAL" else "Draft_Report_REVIEW_REQUIRED.xlsx")
            progress(f"寫入報表：{report_path.name}")
            render(report, tasks, report_path, validation["status"] != "FINAL")
        summary = {"pipeline_version": "5.1", "generated_at": utc_now(), "status": validation["status"], "offline_processing": True, "json_complete_source": True, "source_folder": str(source), "source_document_count": len(files), "processed_document_count": len(documents), "failed_document_count": len(failures), "unaccounted_count": manifest["unaccounted_count"], "item_count": len(shipment["items"]), "review_task_count": len(tasks), "duration_seconds": round(time.time() - started, 2), "outputs": {"report": str(report_path) if report_path else None, "extracted_report": str(output / "extracted_report.json"), "validation": str(output / "validation.json"), "ai_review_request": str(output / "ai_review_request.json"), "document_manifest": str(output / "document_manifest.json"), "normalized_documents": str(output / "normalized")}}
        save_json(output / "processing_summary.json", summary)
        progress(f"完成：狀態={validation['status']}；項目列={len(shipment['items'])}；AI 任務={sum(bool(task.get('ai_editable')) for task in tasks)}")
        return 0 if validation["status"] != "FAILED" else 2
    finally:
        if keep:
            destination = output / "work"
            shutil.rmtree(destination, ignore_errors=True)
            shutil.copytree(workdir, destination)
        shutil.rmtree(workdir, ignore_errors=True)


def parse_date(value: Any, name: str) -> date | None:
    if value is None: return None
    if isinstance(value, datetime): return value.date()
    if isinstance(value, date): return value
    if isinstance(value, str) and value.strip():
        try: return datetime.strptime(value.strip(), "%Y-%m-%d").date()
        except ValueError as exc: raise ValueError(f"{name} 必須使用 YYYY-MM-DD 格式") from exc
    return None


def period(start_value: Any = None, end_value: Any = None) -> tuple[date, date]:
    today = datetime.now(HK_TIMEZONE).date()
    start = parse_date(SEARCH_START_DATE if start_value is None else start_value, "SEARCH_START_DATE")
    end = parse_date(SEARCH_END_DATE if end_value is None else end_value, "SEARCH_END_DATE")
    if start is None and end is None: return today, today
    if start is None: start = end
    if end is None: end = start
    if start is None or end is None or start > end: raise ValueError("日期範圍無效")
    return start, end


def dates(start: date, end: date):
    current = start
    while current <= end:
        yield current
        current += timedelta(days=1)


def discover_shipments(start: date, end: date) -> tuple[list[Path], list[Path]]:
    selected = []
    missing = []
    for day in dates(start, end):
        date_folder = ATTACHMENT_ROOT / day.isoformat()
        progress(f"掃描日期資料夾：{date_folder}")
        if not date_folder.is_dir():
            missing.append(date_folder)
            progress(f"日期資料夾不存在：{date_folder}", "警告")
            continue
        for folder in sorted((path for path in date_folder.iterdir() if path.is_dir() and not path.name.startswith(".")), key=lambda path: path.name.casefold()):
            attachment_folder = folder / ATTACHMENTS_FOLDER
            if attachment_folder.is_dir() and any(path.is_file() and path.suffix.lower() in SUPPORTED_EXTENSIONS for path in attachment_folder.rglob("*")):
                if REPROCESS_EXISTING_REPORTS or not (folder / OUTPUT_FOLDER / "processing_summary.json").is_file(): selected.append(folder)
    return selected, missing


def safe_name(value: str) -> str:
    text = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", norm(value)).rstrip(". ")
    return text[:150] or "NO_SUBJECT"


def unique_path(folder: Path, filename: str) -> Path:
    source = Path(filename)
    destination = folder / f"{safe_name(source.stem)}{source.suffix}"
    counter = 2
    while destination.exists():
        destination = folder / f"{safe_name(source.stem)}_{counter}{source.suffix}"
        counter += 1
    return destination


def subject_reference(subject: str) -> str:
    for values in (contextual_references(subject, MASTER_AIR_LABELS, AIR_MASTER_PATTERN, True), contextual_references(subject, MASTER_SEA_LABELS, SEA_REFERENCE_PATTERN), contextual_references(subject, HOUSE_AIR_LABELS + HOUSE_SEA_LABELS, HOUSE_REFERENCE_PATTERN)):
        if values: return safe_name(values[0])
    return safe_name(subject)


def prepare_msg_shipments(folder: Path) -> list[Path]:
    try:
        import extract_msg
    except ImportError as exc:
        raise RuntimeError("MSG 模式需要 extract-msg") from exc
    if not folder.is_dir(): raise FileNotFoundError(f"MSG 資料夾不存在：{folder}")
    msg_files = sorted((path for path in folder.rglob("*") if path.is_file() and path.suffix.casefold() == ".msg"), key=lambda path: str(path).casefold())
    progress(f"MSG 模式：忽略日期；MSG 數量={len(msg_files)}")
    shipments = []
    for index, msg_path in enumerate(msg_files, start=1):
        message = None
        try:
            progress(f"解析 MSG {index}/{len(msg_files)}：{msg_path.name}")
            message = extract_msg.Message(str(msg_path))
            reference = subject_reference(norm(getattr(message, "subject", "")))
            shipment = ATTACHMENT_ROOT / "MSG_Imported" / reference
            attachments = shipment / ATTACHMENTS_FOLDER
            shipment.mkdir(parents=True, exist_ok=True)
            attachments.mkdir(parents=True, exist_ok=True)
            shutil.copy2(msg_path, unique_path(shipment, msg_path.name))
            count = 0
            for attachment in getattr(message, "attachments", []):
                filename = getattr(attachment, "longFilename", None) or getattr(attachment, "shortFilename", None) or getattr(attachment, "displayName", None) or getattr(attachment, "name", None) or "attachment"
                if Path(filename).suffix.lower() not in SUPPORTED_EXTENSIONS: continue
                destination = unique_path(attachments, filename)
                data = getattr(attachment, "data", None)
                if isinstance(data, bytes): destination.write_bytes(data)
                else: attachment.save(customPath=str(attachments), customFilename=destination.name)
                count += 1
            if count:
                shipments.append(shipment)
                progress(f"MSG 附件已準備：{reference}；有效附件={count}")
            else: progress(f"MSG 沒有 PDF/XLSX：{msg_path.name}", "警告")
        except Exception as exc:
            progress(f"MSG 處理失敗：{msg_path.name}；{type(exc).__name__}: {exc}", "錯誤")
        finally:
            if message is not None:
                try: message.close()
                except Exception: pass
    return list(dict.fromkeys(shipments))


def validate_environment() -> None:
    if not TEMPLATE_PATH.is_file(): raise FileNotFoundError(f"找不到範本：{TEMPLATE_PATH}")
    if not WEIGHT_REFERENCE_PATH.is_file(): raise FileNotFoundError(f"找不到重量參考檔：{WEIGHT_REFERENCE_PATH}")
    if not ATTACHMENT_ROOT.is_dir(): raise FileNotFoundError(f"找不到附件根目錄：{ATTACHMENT_ROOT}")


def run_batch(start: date, end: date, msg_folder: str | None, verbose: bool, keep: bool) -> dict[str, Any]:
    validate_environment()
    if msg_folder:
        folders = prepare_msg_shipments(Path(msg_folder).expanduser().resolve())
        missing = []
    else:
        folders, missing = discover_shipments(start, end)
    result = {"selected": len(folders), "final": 0, "draft": 0, "failed": 0, "missing_dates": [str(path) for path in missing], "shipments": []}
    for index, folder in enumerate(folders, start=1):
        progress(f"總進度 {index}/{len(folders)}：{folder.name}")
        try:
            code = process_shipment(folder, verbose, keep)
            summary_path = folder / OUTPUT_FOLDER / "processing_summary.json"
            status = json.loads(summary_path.read_text(encoding="utf-8")).get("status", "FAILED") if summary_path.is_file() else "FAILED"
            if status == "FINAL": result["final"] += 1
            elif status == "DRAFT_REVIEW_REQUIRED": result["draft"] += 1
            else: result["failed"] += 1
            result["shipments"].append({"folder": str(folder), "status": status, "return_code": code})
        except Exception as exc:
            result["failed"] += 1
            result["shipments"].append({"folder": str(folder), "status": "FAILED", "error": f"{type(exc).__name__}: {exc}"})
            progress(f"批次失敗：{folder.name}；{type(exc).__name__}: {exc}", "錯誤")
    return result


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source")
    parser.add_argument("--start-date")
    parser.add_argument("--end-date")
    parser.add_argument("--msg-folder")
    parser.add_argument("--verbose", action="store_true")
    parser.add_argument("--keep-workdir", action="store_true")
    parser.add_argument("--check", action="store_true")
    args = parser.parse_args()
    try:
        if args.check:
            print(json.dumps({"template_exists": TEMPLATE_PATH.is_file(), "weight_reference_exists": WEIGHT_REFERENCE_PATH.is_file(), "attachment_root_exists": ATTACHMENT_ROOT.is_dir(), "attachment_root": str(ATTACHMENT_ROOT), "default_hong_kong_date": str(datetime.now(HK_TIMEZONE).date()), "configured_msg_folder": MSG_FOLDER_PATH, "dependencies": dependency_status()}, ensure_ascii=False, indent=2))
            return 0
        progress("Step 2 本機文件擷取程序啟動")
        progress("資料處理模式：完全離線；所有 PDF/XLSX 內容只會寫入本機 JSON")
        if args.source: return process_shipment(Path(args.source), args.verbose or VERBOSE_LOGGING, args.keep_workdir or KEEP_WORK_DIRECTORY)
        start, end = period(args.start_date, args.end_date)
        msg_folder = args.msg_folder if args.msg_folder is not None else MSG_FOLDER_PATH
        if msg_folder: progress(f"來源模式：MSG 資料夾；日期設定已忽略；路徑={msg_folder}")
        else: progress(f"來源模式：日期資料夾；香港日期={start} 至 {end}")
        result = run_batch(start, end, msg_folder, args.verbose or VERBOSE_LOGGING, args.keep_workdir or KEEP_WORK_DIRECTORY)
        print("\n" + "=" * 80)
        print("Step 2 技術處理摘要")
        print(f"已選取運輸資料夾：{result['selected']}")
        print(f"最終報表：{result['final']}")
        print(f"待 AI 覆核草稿：{result['draft']}")
        print(f"失敗批次：{result['failed']}")
        print(f"缺少日期資料夾：{len(result['missing_dates'])}")
        print("=" * 80)
        save_json(ASSET_DIR / "step2_last_batch_summary.json", result)
        return 2 if result["failed"] else 0
    except Exception as exc:
        progress(f"程式執行失敗：{type(exc).__name__}: {exc}", "致命錯誤")
        return 2
    finally:
        try: input("\n按 Enter 鍵關閉視窗...")
        except Exception: pass


if __name__ == "__main__":
    raise SystemExit(main())
